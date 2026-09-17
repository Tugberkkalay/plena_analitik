"""
TUSAŞ Excel Data Importer — reads the TUSAS_Plenalitik_Demo_Veri_Seti.xlsx
and converts its three sheets into Plenalitik-compatible MongoDB documents.
"""
import uuid
import random
from datetime import datetime, timezone
from openpyxl import load_workbook

# ── Department name normalization (Excel → taxonomy) ──
_DEPT_MAP = {
    "Montaj Hatti": "Montaj Hattı",
    "CNC ve Talasli Imalat": "CNC ve Talaşlı İmalat",
    "Kompozit Uretim": "Kompozit Üretim",
    "Aviyonik Sistemler": "Aviyonik Sistemler",
    "Yapisal Tasarim": "Yapısal Tasarım",
    "Sistem Muhendisligi": "Sistem Mühendisliği",
    "Yazilim Muhendisligi": "Yazılım Mühendisliği",
    "Test ve Dogrulama": "Test ve Doğrulama",
    "Ucus Bilimleri": "Uçuş Bilimleri",
    "Kalite Guvence": "Kalite Güvence",
    "Tedarik Zinciri": "Tedarik Zinciri",
    "Proje Yonetimi": "Proje Yönetimi",
    "Bilgi Teknolojileri": "Bilgi Teknolojileri",
    "Insan Kaynaklari": "İnsan Kaynakları",
}

_GRUP_MAP = {
    "Uretim": "Üretim",
    "Muhendislik": "Mühendislik",
    "Destek": "Destek",
}

_KADRO_MAP = {
    "Teknisyen": "Teknisyen",
    "Operator": "Operatör",
    "Muhendis": "Mühendis",
    "Uzman": "Uzman",
}

_CINSIYET_MAP = {"Erkek": "Male", "Kadin": "Female"}

_DURUM_MAP = {"Aktif": "active", "Ayrildi": "terminated"}

_CIKIS_MAP = {
    "Gonullu - Yerli Sirket": "Gönüllü - Yerli Şirket",
    "Gonullu - Yurt Disi": "Gönüllü - Yurt Dışı",
    "Gonulsuz": "Gönülsüz",
    "Askerlik/Diger": "Askerlik/Diğer",
    "Emeklilik": "Emeklilik",
}

_KAYNAK_MAP = {
    "Dogrudan Basvuru": "Doğrudan Başvuru",
    "Calisan Referansi": "Çalışan Referansı",
    "SKY Staj": "SKY Programı",
    "LIFT UP": "LIFT UP",
    "MGP": "MGP",
    "Global Talents": "Global Talents",
}

_EGITIM_MAP = {
    "Lisans": "Lisans",
    "Onlisans": "Ön Lisans",
    "Yuksek Lisans": "Yüksek Lisans",
    "Doktora": "Doktora",
}

HRBP_MAP = {
    "Montaj Hattı": "Merve Aksoy", "CNC ve Talaşlı İmalat": "Merve Aksoy", "Kompozit Üretim": "Merve Aksoy",
    "Aviyonik Sistemler": "Elif Korkmaz", "Yapısal Tasarım": "Elif Korkmaz",
    "Sistem Mühendisliği": "Elif Korkmaz", "Yazılım Mühendisliği": "Elif Korkmaz",
    "Test ve Doğrulama": "Deniz Yıldırım", "Uçuş Bilimleri": "Deniz Yıldırım",
    "Kalite Güvence": "Deniz Yıldırım", "Tedarik Zinciri": "Deniz Yıldırım",
    "Proje Yönetimi": "Ayşe Çetin", "Bilgi Teknolojileri": "Ayşe Çetin",
    "İnsan Kaynakları": "Ayşe Çetin",
}

DEPT_SEGMENT_MAP = {
    "Montaj Hattı": "Üretim", "CNC ve Talaşlı İmalat": "Üretim", "Kompozit Üretim": "Üretim",
    "Aviyonik Sistemler": "Mühendislik", "Yapısal Tasarım": "Mühendislik",
    "Sistem Mühendisliği": "Mühendislik", "Yazılım Mühendisliği": "Mühendislik",
    "Test ve Doğrulama": "Mühendislik", "Uçuş Bilimleri": "Mühendislik",
    "Kalite Güvence": "Destek", "Tedarik Zinciri": "Destek",
    "Proje Yönetimi": "Destek", "Bilgi Teknolojileri": "Destek",
    "İnsan Kaynakları": "Destek",
}


def _date_str(dt):
    """Convert datetime to ISO date string."""
    if dt is None:
        return None
    if isinstance(dt, datetime):
        return dt.strftime("%Y-%m-%d")
    return str(dt)


def _salary_to_band(salary):
    """Map salary to band."""
    if salary is None:
        return "B"
    if salary < 55000:
        return "A"
    if salary < 95000:
        return "B"
    if salary < 155000:
        return "C"
    if salary < 240000:
        return "D"
    return "E"


def _read_sheet(wb, sheet_name):
    """Read a sheet into a list of dicts."""
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    return [{headers[i]: row[i] for i in range(min(len(headers), len(row)))} for row in rows[1:]]


def import_tusas_calisanlar(filepath, tenant_slug="tusas"):
    """Import Calisanlar sheet → employees collection format."""
    wb = load_workbook(filepath, read_only=True)
    rows = _read_sheet(wb, "Calisanlar")
    wb.close()

    employees = []
    for r in rows:
        dept = _DEPT_MAP.get(r.get("Bolum", ""), r.get("Bolum", ""))
        salary = r.get("AylikBrutUcret") or 0
        hire_dt = r.get("IseGirisTarihi")
        hire_str = _date_str(hire_dt)
        hire_year = hire_dt.year if isinstance(hire_dt, datetime) else 2024
        seniority = max(0, round(2025 - hire_year + random.uniform(-0.3, 0.3), 1))
        birth_year = r.get("DogumYili") or 1990
        age = 2025 - int(birth_year) if birth_year else 35
        status = _DURUM_MAP.get(r.get("Durum", ""), "active")
        term_dt = r.get("CikisTarihi")
        leave_reason = _CIKIS_MAP.get(r.get("CikisNedeni", ""), r.get("CikisNedeni"))
        term_type = "voluntary" if leave_reason and "Gönüllü" in str(leave_reason) else ("involuntary" if leave_reason else None)
        perf = r.get("PerformansPuani") or 3.5
        band = _salary_to_band(salary)
        is_kritik = str(r.get("KritikRol", "")).lower() in ("evet", "true", "1")

        emp = {
            "id": r.get("SicilNo") or str(uuid.uuid4()),
            "name": r.get("AdSoyad", ""),
            "gender": _CINSIYET_MAP.get(r.get("Cinsiyet", ""), "Male"),
            "age": age,
            "hire_date": hire_str,
            "termination_date": _date_str(term_dt),
            "department": dept,
            "job_title": _KADRO_MAP.get(r.get("Kadro", ""), r.get("Kadro", "")),
            "band": band,
            "salary": int(salary) if salary else 50000,
            "city": r.get("Lokasyon", "Ankara Kahramankazan"),
            "country": "Turkey",
            "education_level": _EGITIM_MAP.get(r.get("EgitimSeviyesi", ""), r.get("EgitimSeviyesi", "")),
            "university": r.get("Universite"),
            "marital_status": random.choice(["Bekar", "Evli"]),
            "is_talent": is_kritik or (perf >= 4.3),
            "is_manager": band in ("D", "E"),
            "is_disabled": False,
            "is_full_time": True,
            "performance_score": round(float(perf), 1) if perf else 3.5,
            "seniority_years": seniority,
            "mobility_flag": random.random() < 0.25,
            "branch_id": "",
            "region": r.get("Lokasyon", ""),
            "role_type": "core",
            "segment": _GRUP_MAP.get(r.get("Grup", ""), r.get("Grup", "")),
            "hrbp": HRBP_MAP.get(dept, ""),
            "status": status,
            "leaving_reason": leave_reason,
            "termination_type": term_type,
            "project": r.get("Proje", ""),
            "recruitment_source": _KAYNAK_MAP.get(r.get("IseAlimKaynagi", ""), r.get("IseAlimKaynagi", "")),
            "english_level": r.get("IngilizceSeviyesi", ""),
            "project_bonus": str(r.get("ProjePrimiAliyorMu", "")).lower() in ("evet", "true", "1"),
            "engagement_score": r.get("BaglilikSkoru") or 70,
            "data_source": "excel_import",
            "tenant_id": tenant_slug,
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        employees.append(emp)
    return employees


def import_tusas_recruitment(filepath, tenant_slug="tusas"):
    """Import IseAlimHunisi sheet → recruitment collection format."""
    wb = load_workbook(filepath, read_only=True)
    rows = _read_sheet(wb, "IseAlimHunisi")
    wb.close()

    STAGE_ORDER = ["Basvuru", "Online Degerlendirme", "IK Mulakati", "Teknik Mulakat",
                   "Guvenlik Sorusturmasi", "Teklif", "Ise Baslama"]
    STAGE_MAP = {
        "Basvuru": "Başvuru", "Online Degerlendirme": "Online Değerlendirme",
        "IK Mulakati": "İK Mülakatı", "Teknik Mulakat": "Teknik Mülakat",
        "Guvenlik Sorusturmasi": "Güvenlik Soruşturması", "Teklif": "Teklif",
        "Ise Baslama": "İşe Başlama",
    }
    DURUM_MAP_R = {
        "Ise Baslatildi": "İşe Başlatıldı",
        "Elendi/Vazgecti": "Elendi/Vazgeçti",
        "Surecte": "Süreçte",
    }

    records = []
    for r in rows:
        dept_raw = r.get("Bolum", "")
        dept = _DEPT_MAP.get(dept_raw, dept_raw)
        son_asama = r.get("SonAsama", "")
        durum = r.get("Durum", "")
        # Determine funnel stage index
        stage_idx = 0
        for i, s in enumerate(STAGE_ORDER):
            if son_asama and s.lower().replace(" ", "") in son_asama.lower().replace(" ", ""):
                stage_idx = i
                break
        hired = "Ise Baslatildi" in str(durum)

        rec = {
            "id": r.get("BasvuruNo") or str(uuid.uuid4()),
            "name": f"Aday-{r.get('BasvuruNo', '')}",
            "gender": random.choice(["Male", "Female"]),
            "department": dept,
            "position": r.get("Pozisyon", ""),
            "band": "B" if "Muhendis" in str(r.get("Kadro", "")) else "A",
            "source": r.get("BasvuruKanali", ""),
            "application_date": _date_str(r.get("BasvuruTarihi")),
            "stage": STAGE_MAP.get(son_asama, son_asama),
            "stage_index": stage_idx,
            "status": DURUM_MAP_R.get(durum, durum),
            "hired": hired,
            "online_eval_date": _date_str(r.get("OnlineDegTarihi")),
            "hr_interview_date": _date_str(r.get("IKMulakatTarihi")),
            "technical_interview_date": _date_str(r.get("TeknikMulakatTarihi")),
            "security_clearance_date": _date_str(r.get("GuvenlikSorusturmaTamamTarihi")),
            "offer_date": _date_str(r.get("TeklifTarihi")),
            "start_date": _date_str(r.get("IseBaslamaTarihi")),
            "total_days": r.get("ToplamSurecGun") or 0,
            "university": random.choice(["ODTÜ", "İTÜ", "Hacettepe", "Bilkent", "Yıldız Teknik", "Eskişehir Teknik", "Gazi"]),
            "education": random.choice(["Lisans", "Yüksek Lisans"]),
            "rejection_reason": "" if hired else random.choice(["Teknik Yeterlilik", "Güvenlik Soruşturması", "Aday Vazgeçti", "Başka Teklif", "Kontenjan Doldu", ""]),
            "segment": DEPT_SEGMENT_MAP.get(dept, ""),
            "hrbp": HRBP_MAP.get(dept, ""),
            "tenant_id": tenant_slug,
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        records.append(rec)
    return records


def import_tusas_yetenek(filepath, tenant_slug="tusas"):
    """Import YetenekProgramlari sheet → talent_programs collection."""
    wb = load_workbook(filepath, read_only=True)
    rows = _read_sheet(wb, "YetenekProgramlari")
    wb.close()

    records = []
    for r in rows:
        records.append({
            "id": r.get("KatilimciNo") or str(uuid.uuid4()),
            "program": r.get("Program", ""),
            "donem": r.get("Donem"),
            "cinsiyet": _CINSIYET_MAP.get(r.get("Cinsiyet", ""), "Male"),
            "universite": r.get("Universite", ""),
            "tamamladi": str(r.get("ProgramiTamamladi", "")).lower() in ("evet", "true", "1"),
            "ise_alindi": str(r.get("IseAlindi", "")).lower() in ("evet", "true", "1"),
            "ilk_yil_kaldi": str(r.get("Ilk1YilKaldiMi", "")).lower() in ("evet", "true", "1"),
            "tenant_id": tenant_slug,
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
    return records
