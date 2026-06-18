"""Excel template generation and data import for Plenalitik."""
import io
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, timezone

HEADER_FILL = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
CELL_FONT = Font(name="Calibri", size=10)
THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"), right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"), bottom=Side(style="thin", color="D1D5DB")
)

SHEETS = {
    "Çalışanlar": {
        "columns": [
            ("Çalışan ID", 15), ("Ad Soyad", 25), ("Cinsiyet", 10), ("Yaş", 8), ("İşe Giriş Tarihi", 18),
            ("Ayrılma Tarihi", 18), ("Departman", 22), ("Pozisyon", 30), ("Band (A-E)", 12),
            ("Maaş", 15), ("Şehir", 15), ("Ülke", 12), ("Eğitim Düzeyi", 18), ("Üniversite", 25),
            ("Medeni Durum", 15), ("Yetenek mi?", 12), ("Yönetici mi?", 12), ("Engelli mi?", 12),
            ("Tam Zamanlı mı?", 14), ("Performans (1-5)", 14), ("Kıdem (Yıl)", 12),
            ("Şube ID", 15), ("Bölge", 15), ("Durum", 12), ("Ayrılma Nedeni", 20),
        ],
        "sample": [
            "E001", "Ahmet Yılmaz", "Male", 35, "2020-03-15", "", "Kredi ve Risk", "Kredi Analisti", "C",
            85000, "İstanbul", "Turkey", "Lisans", "İstanbul Üniversitesi", "Evli", "Evet", "Hayır", "Hayır",
            "Evet", 4.2, 5, "BR001", "Marmara", "active", "",
        ],
    },
    "Yetkinlikler": {
        "columns": [
            ("Çalışan ID", 15), ("Yetkinlik Adı", 35), ("Seviye (1-5)", 12),
        ],
        "sample": ["E001", "Basel III/IV Uygulamaları", 4],
    },
    "Şubeler": {
        "columns": [
            ("Şube ID", 15), ("Şube Adı", 25), ("Bölge", 15), ("Şehir", 15),
            ("Enlem", 12), ("Boylam", 12), ("Segment", 15), ("Hedef Kadro", 12),
        ],
        "sample": ["BR001", "Levent Şubesi", "Marmara", "İstanbul", 41.08, 29.01, "Bireysel", 25],
    },
    "Satış Performansı": {
        "columns": [
            ("Çalışan ID", 15), ("Şube ID", 15), ("Dönem (YYYY-MM)", 18),
            ("Hedef", 15), ("Gerçekleşen", 15), ("Portföy Büyüklüğü", 18),
        ],
        "sample": ["E001", "BR001", "2025-01", 500000, 580000, 2500000],
    },
    "İşe Alım": {
        "columns": [
            ("Pozisyon", 30), ("Departman", 22), ("Açılma Tarihi", 18), ("Kapanma Tarihi", 18),
            ("Başvuru Sayısı", 14), ("Mülakat Sayısı", 14), ("Teklif Sayısı", 12),
            ("Durum", 15), ("Kaynak", 18), ("İşe Alım Süresi (Gün)", 20),
        ],
        "sample": [
            "Kredi Analisti", "Kredi ve Risk", "2025-01-10", "2025-02-15", 45, 8, 2,
            "Closed", "LinkedIn", 36,
        ],
    },
    "Eğitimler": {
        "columns": [
            ("Çalışan ID", 15), ("Eğitim Adı", 35), ("Kategori", 20), ("Tamamlanma Tarihi", 18),
            ("Süre (Saat)", 12), ("Puan (0-100)", 12), ("Durum", 15),
        ],
        "sample": ["E001", "Basel IV Temel Eğitim", "Teknik", "2025-03-20", 16, 85, "Tamamlandı"],
    },
    "Bağlılık Anketi": {
        "columns": [
            ("Çalışan ID", 15), ("Genel Puan (1-5)", 14), ("İş Tatmini (1-5)", 14),
            ("Yönetim (1-5)", 14), ("Kariyer Gelişimi (1-5)", 14), ("İş-Yaşam Dengesi (1-5)", 16),
            ("Tavsiye Eder mi? (1-10)", 18),
        ],
        "sample": ["E001", 4.2, 4.0, 4.5, 3.8, 3.5, 8],
    },
}


def generate_template() -> bytes:
    """Generate an empty Excel template with all sheets and headers."""
    wb = Workbook()
    first = True
    for sheet_name, config in SHEETS.items():
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(sheet_name)

        # Write headers
        for col_idx, (col_name, width) in enumerate(config["columns"], 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER
            ws.column_dimensions[cell.column_letter].width = width

        # Write sample row
        sample = config["sample"]
        for col_idx, val in enumerate(sample, 1):
            cell = ws.cell(row=2, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10, italic=True, color="999999")
            cell.border = THIN_BORDER

        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_excel(file_bytes: bytes, tenant_slug: str) -> dict:
    """Parse uploaded Excel and return structured data for DB insertion."""
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    result = {"employees": [], "skills_map": {}, "branches": [], "sales": [], "recruitment": [], "training": [], "engagement": []}

    # 1. Çalışanlar
    if "Çalışanlar" in wb.sheetnames:
        ws = wb["Çalışanlar"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0] or not row[1]:
                continue
            emp = {
                "id": str(row[0]).strip(),
                "name": str(row[1]).strip(),
                "gender": str(row[2] or "Male").strip(),
                "age": int(row[3] or 30),
                "hire_date": str(row[4] or "2024-01-01").strip()[:10],
                "termination_date": str(row[5]).strip()[:10] if row[5] else None,
                "department": str(row[6] or "").strip(),
                "job_title": str(row[7] or "").strip(),
                "band": str(row[8] or "B").strip().upper()[:1],
                "salary": float(row[9] or 50000),
                "city": str(row[10] or "İstanbul").strip(),
                "country": str(row[11] or "Turkey").strip(),
                "education_level": str(row[12] or "Lisans").strip(),
                "university": str(row[13] or "").strip(),
                "marital_status": str(row[14] or "Bekar").strip(),
                "is_talent": str(row[15] or "").strip().lower() in ("evet", "yes", "true", "1"),
                "is_manager": str(row[16] or "").strip().lower() in ("evet", "yes", "true", "1"),
                "is_disabled": str(row[17] or "").strip().lower() in ("evet", "yes", "true", "1"),
                "is_full_time": str(row[18] or "Evet").strip().lower() not in ("hayır", "no", "false", "0"),
                "performance_score": min(5, max(0, float(row[19] or 3))),
                "seniority_years": float(row[20] or 1),
                "branch_id": str(row[21] or "").strip(),
                "region": str(row[22] or "").strip(),
                "status": str(row[23] or "active").strip().lower(),
                "leaving_reason": str(row[24] or "").strip() if row[24] else None,
                "tenant_id": tenant_slug,
                "skills": [],
                "mobility_flag": False,
                "role_type": "specialist",
                "data_source": "Excel Import",
                "created_at": datetime.now(timezone.utc).isoformat(),
            }
            if emp["is_manager"]:
                emp["role_type"] = "manager"
            result["employees"].append(emp)

    # 2. Yetkinlikler
    if "Yetkinlikler" in wb.sheetnames:
        ws = wb["Yetkinlikler"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0] or not row[1]:
                continue
            emp_id = str(row[0]).strip()
            if emp_id not in result["skills_map"]:
                result["skills_map"][emp_id] = []
            result["skills_map"][emp_id].append({
                "skill": str(row[1]).strip(),
                "proficiency": min(5, max(1, int(row[2] or 3))),
            })

    # Apply skills to employees
    for emp in result["employees"]:
        emp["skills"] = result["skills_map"].get(emp["id"], [])

    # 3. Şubeler
    if "Şubeler" in wb.sheetnames:
        ws = wb["Şubeler"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0] or not row[1]:
                continue
            result["branches"].append({
                "id": str(row[0]).strip(),
                "name": str(row[1]).strip(),
                "region": str(row[2] or "").strip(),
                "city": str(row[3] or "").strip(),
                "lat": float(row[4] or 0),
                "lng": float(row[5] or 0),
                "segment": str(row[6] or "Bireysel").strip(),
                "target_headcount": int(row[7] or 20),
                "headcount": 0,
                "tenant_id": tenant_slug,
            })

    # Calculate branch headcounts
    branch_counts = {}
    for emp in result["employees"]:
        if emp["status"] == "active" and emp["branch_id"]:
            branch_counts[emp["branch_id"]] = branch_counts.get(emp["branch_id"], 0) + 1
    for br in result["branches"]:
        br["headcount"] = branch_counts.get(br["id"], 0)

    # 4. Satış Performansı
    if "Satış Performansı" in wb.sheetnames:
        ws = wb["Satış Performansı"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            target = float(row[3] or 0)
            actual = float(row[4] or 0)
            ach = round(actual / target * 100, 1) if target else 0
            # Commission tiers
            if ach < 85: comm_rate = 0
            elif ach < 100: comm_rate = 0.005
            elif ach < 120: comm_rate = 0.01
            else: comm_rate = 0.015
            result["sales"].append({
                "employee_id": str(row[0]).strip(),
                "branch_id": str(row[1] or "").strip(),
                "period": str(row[2] or "2025-01").strip(),
                "target": target, "actual": actual,
                "achievement_pct": ach,
                "commission": round(actual * comm_rate),
                "portfolio_size": float(row[5] or 0),
                "tenant_id": tenant_slug,
            })

    # 5. İşe Alım
    if "İşe Alım" in wb.sheetnames:
        ws = wb["İşe Alım"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            result["recruitment"].append({
                "position": str(row[0]).strip(),
                "department": str(row[1] or "").strip(),
                "opened_date": str(row[2] or "").strip()[:10],
                "closed_date": str(row[3] or "").strip()[:10] if row[3] else None,
                "applications": int(row[4] or 0),
                "interviews": int(row[5] or 0),
                "offers": int(row[6] or 0),
                "status": str(row[7] or "Open").strip(),
                "source": str(row[8] or "").strip(),
                "time_to_fill": int(row[9] or 0) if row[9] else None,
                "tenant_id": tenant_slug,
            })

    # 6. Eğitimler
    if "Eğitimler" in wb.sheetnames:
        ws = wb["Eğitimler"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0] or not row[1]:
                continue
            result["training"].append({
                "employee_id": str(row[0]).strip(),
                "course_name": str(row[1]).strip(),
                "category": str(row[2] or "Genel").strip(),
                "completion_date": str(row[3] or "").strip()[:10] if row[3] else None,
                "duration_hours": float(row[4] or 0),
                "score": float(row[5] or 0) if row[5] else None,
                "status": str(row[6] or "Tamamlandı").strip(),
                "tenant_id": tenant_slug,
            })

    # 7. Bağlılık Anketi
    if "Bağlılık Anketi" in wb.sheetnames:
        ws = wb["Bağlılık Anketi"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            result["engagement"].append({
                "employee_id": str(row[0]).strip(),
                "overall_score": float(row[1] or 3),
                "job_satisfaction": float(row[2] or 3),
                "management": float(row[3] or 3),
                "career_development": float(row[4] or 3),
                "work_life_balance": float(row[5] or 3),
                "nps_score": int(row[6] or 7) if row[6] else 7,
                "tenant_id": tenant_slug,
            })

    del result["skills_map"]
    return result
