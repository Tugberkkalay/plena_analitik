"""
Jolly Tur Excel Importer & Performance Analytics Endpoints.
Imports 5 sheets: Parametreler, Calisanlar, Hedefler, Yetkinlikler, Karne_Ozeti.
Provides 6 dashboard screens: Overview, Target, Competency, Calibration, Talent Matrix, Employee Card.
"""
import random
from datetime import datetime, timezone
from collections import Counter, defaultdict
from openpyxl import load_workbook
from fastapi import APIRouter, HTTPException

router = APIRouter(prefix="/api/jolly", tags=["jolly-performance"])
db = None

def setup_jolly(database):
    global db
    db = database


def _read_named_sheet(wb, sheet_name):
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    return [{headers[i]: row[i] for i in range(min(len(headers), len(row)))} for row in rows[1:] if any(c is not None for c in row)]


def _safe_float(val, default=0.0):
    if val is None:
        return default
    try:
        return float(val)
    except (ValueError, TypeError):
        return default


async def import_jolly_data(filepath, tenant_slug="jollytur"):
    """Import all 5 sheets from Jolly Tur Excel file."""
    wb = load_workbook(filepath, data_only=True)

    # 1. Parametreler - special structure
    ws_param = wb["Parametreler"]
    param_rows = list(ws_param.iter_rows(values_only=True))
    param_docs = []
    segment_defs = []
    for row in param_rows:
        if row[0] and row[1] is not None:
            if str(row[0]).strip() in ("Satış Danışmanı", "Çağrı Merkezi Temsilcisi", "Şube Müdürü", "Çağrı Merkezi Yöneticisi", "Satış Direktörü"):
                param_docs.append({
                    "rol": str(row[0]).strip(),
                    "hedef_agirligi": _safe_float(row[1]),
                    "yetkinlik_agirligi": _safe_float(row[2]),
                    "tenant_id": tenant_slug,
                })
            elif str(row[0]).strip() in ("Yıldız", "Güçlü", "Beklenen", "Gelişim Alanı"):
                segment_defs.append({
                    "segment": str(row[0]).strip(),
                    "alt_sinir": _safe_float(row[1]),
                    "aciklama": str(row[2] or ""),
                    "tenant_id": tenant_slug,
                })

    # 2. Calisanlar
    calisanlar = _read_named_sheet(wb, "Calisanlar")

    # 3. Hedefler
    hedefler_raw = _read_named_sheet(wb, "Hedefler")

    # 4. Yetkinlikler
    yetkinlikler_raw = _read_named_sheet(wb, "Yetkinlikler")

    # 5. Karne_Ozeti
    karne_raw = _read_named_sheet(wb, "Karne_Ozeti")
    wb.close()

    # Clear existing data
    for coll in ["jt_parametreler", "jt_hedefler", "jt_yetkinlikler", "jt_karne"]:
        await db[coll].delete_many({"tenant_id": tenant_slug})
    await db.employees.delete_many({"tenant_id": tenant_slug})

    # Insert Parametreler
    if param_docs:
        await db.jt_parametreler.insert_many(param_docs)

    # Build karne map for employee enrichment
    karne_map = {}
    for k in karne_raw:
        sicil = str(k.get("Sicil No", ""))
        karne_map[sicil] = k

    # Insert employees
    emp_docs = []
    for c in calisanlar:
        sicil = str(c.get("Sicil No", ""))
        k = karne_map.get(sicil, {})
        rol = c.get("Rol", "")
        emp_docs.append({
            "id": sicil,
            "name": c.get("Ad Soyad", ""),
            "job_title": rol,
            "department": c.get("Şube", ""),
            "city": c.get("Bölge", ""),
            "segment": k.get("Segment", ""),
            "hrbp": "",
            "band": "C" if "Müdür" in str(rol) or "Direktör" in str(rol) else "B",
            "salary": random.randint(30000, 80000),
            "gender": random.choice(["Male", "Female"]),
            "age": random.randint(25, 55),
            "hire_date": str(c.get("İşe Giriş Tarihi", ""))[:10] if c.get("İşe Giriş Tarihi") else "2023-01-15",
            "status": "active",
            "performance_score": round(_safe_float(k.get("Toplam Skor")) / 25, 1),
            "manager_id": str(c.get("Yönetici Sicil", "") or ""),
            "hedef_skoru": round(_safe_float(k.get("Hedef Skoru")), 2),
            "yetkinlik_skoru": round(_safe_float(k.get("Yetkinlik Skoru")), 2),
            "toplam_skor": round(_safe_float(k.get("Toplam Skor")), 2),
            "oz_yonetici_farki": round(_safe_float(k.get("Öz-Yönetici Ort. Farkı")), 2),
            "hedef_agirligi": _safe_float(k.get("Hedef Ağırlığı")),
            "yetkinlik_agirligi": _safe_float(k.get("Yetkinlik Ağırlığı")),
            "country": "Turkey",
            "region": c.get("Bölge", ""),
            "is_talent": k.get("Segment") == "Yıldız",
            "is_manager": "Müdür" in str(rol) or "Direktör" in str(rol),
            "skills": [],
            "tenant_id": tenant_slug,
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
    if emp_docs:
        await db.employees.insert_many(emp_docs)

    # Insert Hedefler
    hedef_docs = []
    for h in hedefler_raw:
        hedef_docs.append({
            "sicil_no": str(h.get("Sicil No", "")),
            "ad_soyad": h.get("Ad Soyad", ""),
            "rol": h.get("Rol", ""),
            "sube": h.get("Şube", ""),
            "bolge": h.get("Bölge", ""),
            "hedef": h.get("Hedef", ""),
            "agirlik": _safe_float(h.get("Ağırlık (%)")),
            "birim": h.get("Birim", ""),
            "hedef_deger": _safe_float(h.get("Hedef Değer")),
            "gerceklesen": _safe_float(h.get("Gerçekleşen")),
            "gerceklesme_pct": min(120, round(_safe_float(h.get("Gerçekleşme (%)")), 2)),
            "agirlikli_skor": round(_safe_float(h.get("Ağırlıklı Skor")), 2),
            "tenant_id": tenant_slug,
        })
    if hedef_docs:
        await db.jt_hedefler.insert_many(hedef_docs)

    # Insert Yetkinlikler
    yetkinlik_docs = []
    for y in yetkinlikler_raw:
        yetkinlik_docs.append({
            "sicil_no": str(y.get("Sicil No", "")),
            "ad_soyad": y.get("Ad Soyad", ""),
            "rol": y.get("Rol", ""),
            "sube": y.get("Şube", ""),
            "bolge": y.get("Bölge", ""),
            "yetkinlik_no": y.get("Yetkinlik No", 0),
            "yetkinlik": y.get("Yetkinlik", ""),
            "kategori": y.get("Kategori", ""),
            "oz_degerlendirme": _safe_float(y.get("Öz Değerlendirme (1-5)")),
            "yonetici_puani": _safe_float(y.get("Yönetici Puanı (1-5)")),
            "fark": _safe_float(y.get("Fark (Öz - Yönetici)")),
            "tenant_id": tenant_slug,
        })
    if yetkinlik_docs:
        await db.jt_yetkinlikler.insert_many(yetkinlik_docs)

    # Insert Karne Özeti
    karne_docs = []
    for k in karne_raw:
        karne_docs.append({
            "sicil_no": str(k.get("Sicil No", "")),
            "ad_soyad": k.get("Ad Soyad", ""),
            "rol": k.get("Rol", ""),
            "sube": k.get("Şube", ""),
            "bolge": k.get("Bölge", ""),
            "hedef_agirligi": _safe_float(k.get("Hedef Ağırlığı")),
            "yetkinlik_agirligi": _safe_float(k.get("Yetkinlik Ağırlığı")),
            "hedef_skoru": round(_safe_float(k.get("Hedef Skoru")), 2),
            "yetkinlik_skoru": round(_safe_float(k.get("Yetkinlik Skoru")), 2),
            "toplam_skor": round(_safe_float(k.get("Toplam Skor")), 2),
            "segment": k.get("Segment", ""),
            "oz_yonetici_farki": round(_safe_float(k.get("Öz-Yönetici Ort. Farkı")), 2),
            "tenant_id": tenant_slug,
        })
    if karne_docs:
        await db.jt_karne.insert_many(karne_docs)

    return {
        "parametreler": len(param_docs), "calisanlar": len(emp_docs),
        "hedefler": len(hedef_docs), "yetkinlikler": len(yetkinlik_docs),
        "karne": len(karne_docs),
    }


# ─── Analytics Endpoints ───

def _build_filter(tenant, bolge=None, sube=None, rol=None):
    q = {"tenant_id": tenant} if tenant else {}
    if bolge:
        q["bolge"] = bolge
    if sube:
        q["sube"] = sube
    if rol:
        q["rol"] = rol
    return q


@router.get("/overview")
async def jt_overview(tenant: str = None, bolge: str = None, sube: str = None, rol: str = None):
    """Screen 1: Genel Bakış"""
    q = _build_filter(tenant, bolge, sube, rol)
    karne = await db.jt_karne.find(q, {"_id": 0}).to_list(500)
    if not karne:
        return {"kpis": {}, "segment_dist": [], "skor_histogram": [], "bolge_skor": []}

    n = len(karne)
    avg_toplam = round(sum(k["toplam_skor"] for k in karne) / n, 2)
    avg_hedef = round(sum(k["hedef_skoru"] for k in karne) / n, 2)
    avg_yetkinlik = round(sum(k["yetkinlik_skoru"] for k in karne) / n, 2)
    seg_counts = Counter(k["segment"] for k in karne)
    yildiz_pct = round(seg_counts.get("Yıldız", 0) / n * 100, 1)

    seg_colors = {"Yıldız": "#F59E0B", "Güçlü": "#14B8A6", "Beklenen": "#6366F1", "Gelişim Alanı": "#EF4444"}
    segment_dist = [{"name": s, "value": c, "color": seg_colors.get(s, "#94a3b8")} for s, c in seg_counts.most_common()]

    buckets = [(f"{lo}-{lo+10}", lo, lo+10) for lo in range(40, 120, 10)]
    skor_hist = [{"range": label, "count": len([k for k in karne if lo <= k["toplam_skor"] < hi])} for label, lo, hi in buckets]

    bolge_groups = defaultdict(list)
    for k in karne:
        bolge_groups[k.get("bolge", "")].append(k["toplam_skor"])
    bolge_skor = sorted([{"bolge": b, "ort_skor": round(sum(v)/len(v), 1), "kisi": len(v)} for b, v in bolge_groups.items()], key=lambda x: -x["ort_skor"])

    return {
        "kpis": {"calisan_sayisi": n, "ort_toplam_skor": avg_toplam, "ort_hedef_gerceklesme": avg_hedef,
                 "ort_yetkinlik_skoru": avg_yetkinlik, "yildiz_orani": yildiz_pct},
        "segment_dist": segment_dist, "skor_histogram": skor_hist, "bolge_skor": bolge_skor,
    }


@router.get("/hedef-analizi")
async def jt_hedef_analizi(tenant: str = None, bolge: str = None, sube: str = None, rol: str = None):
    """Screen 2: Hedef Analizi"""
    q = _build_filter(tenant, bolge, sube, rol)
    hedefler = await db.jt_hedefler.find(q, {"_id": 0}).to_list(5000)
    karne = await db.jt_karne.find(q, {"_id": 0}).to_list(500)
    if not hedefler:
        return {"hedef_ort": [], "bolge_hedef_heatmap": [], "sube_siralama": [], "scatter": []}

    # Target-level average completion
    hedef_groups = defaultdict(list)
    for h in hedefler:
        name = h.get("hedef", "")
        if name:
            hedef_groups[name].append(h["gerceklesme_pct"])
    hedef_ort = sorted([{"hedef": h, "ort_gerceklesme": round(sum(v)/len(v), 1)} for h, v in hedef_groups.items()], key=lambda x: x["ort_gerceklesme"])

    # Region x Target heatmap
    bolge_hedef = defaultdict(lambda: defaultdict(list))
    for h in hedefler:
        name = h.get("hedef", "")
        if name:
            bolge_hedef[h.get("bolge", "")][name].append(h["gerceklesme_pct"])
    heatmap = []
    for b, targets in bolge_hedef.items():
        for t, vals in targets.items():
            heatmap.append({"bolge": b, "hedef": t, "ort": round(sum(vals)/len(vals), 1)})

    # Branch ranking by target score
    sube_groups = defaultdict(list)
    for k in karne:
        sube_groups[k.get("sube", "")].append(k["hedef_skoru"])
    sube_sira = sorted([{"sube": s, "ort_hedef": round(sum(v)/len(v), 1), "kisi": len(v)} for s, v in sube_groups.items()], key=lambda x: -x["ort_hedef"])

    # Scatter: find Ciro target realization vs NPS/müşteri memnuniyeti per person
    scatter = []
    kisi_ciro = {}
    kisi_nps = {}
    for h in hedefler:
        sicil = h.get("sicil_no", "")
        hedef_name = str(h.get("hedef", "")).lower()
        birim = str(h.get("birim", "")).lower()
        if "ciro" in hedef_name or "satış ciro" in birim:
            kisi_ciro[sicil] = h["gerceklesme_pct"]
        if "nps" in hedef_name or "memnuniyet" in hedef_name or "nps" in birim:
            kisi_nps[sicil] = h["gerceklesme_pct"]
    for sicil in kisi_ciro:
        if sicil in kisi_nps:
            k = next((x for x in karne if x.get("sicil_no") == sicil), {})
            scatter.append({"ad": k.get("ad_soyad", sicil), "ciro_gerceklesme": round(kisi_ciro[sicil], 1),
                           "nps_gerceklesme": round(kisi_nps[sicil], 1), "sube": k.get("sube", "")})

    return {"hedef_ort": hedef_ort, "bolge_hedef_heatmap": heatmap, "sube_siralama": sube_sira, "scatter": scatter}


@router.get("/yetkinlik-analizi")
async def jt_yetkinlik_analizi(tenant: str = None, bolge: str = None, sube: str = None, rol: str = None):
    """Screen 3: Yetkinlik Analizi"""
    q = _build_filter(tenant, bolge, sube, rol)
    yetkinlikler = await db.jt_yetkinlikler.find(q, {"_id": 0}).to_list(5000)
    if not yetkinlikler:
        return {"kategori_radar": [], "yetkinlik_sube_heatmap": [], "en_guclu": [], "en_zayif": [], "rol_karsilastirma": []}

    kat_groups = defaultdict(list)
    for y in yetkinlikler:
        kat_groups[y.get("kategori", "")].append(y["yonetici_puani"])
    kategori_radar = [{"kategori": k, "ort_puan": round(sum(v)/len(v), 2)} for k, v in kat_groups.items() if k]

    yetkinlik_sube = defaultdict(lambda: defaultdict(list))
    for y in yetkinlikler:
        yname = y.get("yetkinlik", "")
        if yname:
            yetkinlik_sube[yname][y.get("sube", "")].append(y["yonetici_puani"])
    heatmap = []
    for yname, subs in yetkinlik_sube.items():
        for s, vals in subs.items():
            heatmap.append({"yetkinlik": yname, "sube": s, "ort_puan": round(sum(vals)/len(vals), 2)})

    yetkinlik_avg = defaultdict(list)
    for y in yetkinlikler:
        yname = y.get("yetkinlik", "")
        if yname:
            yetkinlik_avg[yname].append(y["yonetici_puani"])
    all_avg = sorted([{"yetkinlik": y, "ort_puan": round(sum(v)/len(v), 2)} for y, v in yetkinlik_avg.items()], key=lambda x: -x["ort_puan"])
    en_guclu = all_avg[:3]
    en_zayif = all_avg[-3:]

    rol_groups = defaultdict(list)
    for y in yetkinlikler:
        rol_groups[y.get("rol", "")].append(y["yonetici_puani"])
    rol_karsilastirma = [{"rol": r, "ort_puan": round(sum(v)/len(v), 2)} for r, v in rol_groups.items()]

    return {"kategori_radar": kategori_radar, "yetkinlik_sube_heatmap": heatmap,
            "en_guclu": en_guclu, "en_zayif": en_zayif, "rol_karsilastirma": rol_karsilastirma}


@router.get("/kalibrasyon")
async def jt_kalibrasyon(tenant: str = None, bolge: str = None, sube: str = None, rol: str = None):
    """Screen 4: Kalibrasyon ve Değerlendirici Analizi"""
    q = _build_filter(tenant, bolge, sube, rol)
    yetkinlikler = await db.jt_yetkinlikler.find(q, {"_id": 0}).to_list(5000)
    karne = await db.jt_karne.find(q, {"_id": 0}).to_list(500)
    if not yetkinlikler:
        return {"sube_karsilastirma": [], "degerlendirici_sapma": [], "fark_top10": []}

    # Branch: self vs manager average
    sube_oz = defaultdict(list)
    sube_yon = defaultdict(list)
    for y in yetkinlikler:
        sube_oz[y.get("sube", "")].append(y["oz_degerlendirme"])
        sube_yon[y.get("sube", "")].append(y["yonetici_puani"])
    sube_karsilastirma = [{"sube": s, "oz_ort": round(sum(sube_oz[s])/len(sube_oz[s]), 2),
                          "yonetici_ort": round(sum(sube_yon[s])/len(sube_yon[s]), 2)} for s in sube_oz]

    # Evaluator severity analysis
    company_avg = sum(y["yonetici_puani"] for y in yetkinlikler) / len(yetkinlikler) if yetkinlikler else 0
    emp_q = {"tenant_id": q.get("tenant_id", "")} if q.get("tenant_id") else {}
    all_emps = await db.employees.find(emp_q, {"_id": 0, "id": 1, "name": 1, "manager_id": 1}).to_list(200)
    emp_mgr = {e["id"]: e.get("manager_id", "") for e in all_emps}
    mgr_names = {e["id"]: e["name"] for e in all_emps}

    mgr_scores = defaultdict(list)
    mgr_people = defaultdict(set)
    for y in yetkinlikler:
        sicil = y.get("sicil_no", "")
        mgr_id = emp_mgr.get(sicil, "")
        if mgr_id:
            mgr_scores[mgr_id].append(y["yonetici_puani"])
            mgr_people[mgr_id].add(sicil)

    degerlendirici_sapma = []
    for mgr_id, scores in mgr_scores.items():
        avg = sum(scores) / len(scores)
        sapma = round(avg - company_avg, 2)
        tip = "Cömert" if sapma > 0.3 else ("Sert" if sapma < -0.3 else "Normal")
        degerlendirici_sapma.append({
            "yonetici": mgr_names.get(mgr_id, mgr_id), "sicil": mgr_id,
            "ort_puan": round(avg, 2), "sirket_ort": round(company_avg, 2),
            "sapma": sapma, "tip": tip, "degerlendirilen": len(mgr_people[mgr_id])
        })
    degerlendirici_sapma.sort(key=lambda x: x["sapma"])

    # Top 10 highest self-manager gap
    kisi_fark = defaultdict(list)
    for y in yetkinlikler:
        kisi_fark[y["sicil_no"]].append(y["fark"])
    fark_top10 = []
    for sicil, farks in kisi_fark.items():
        avg_fark = sum(farks) / len(farks)
        k = next((x for x in karne if x["sicil_no"] == sicil), {})
        fark_top10.append({"sicil": sicil, "ad": k.get("ad_soyad", ""), "sube": k.get("sube", ""),
                          "ort_fark": round(avg_fark, 2), "segment": k.get("segment", "")})
    fark_top10.sort(key=lambda x: abs(x["ort_fark"]), reverse=True)

    return {"sube_karsilastirma": sube_karsilastirma, "degerlendirici_sapma": degerlendirici_sapma, "fark_top10": fark_top10[:10]}


@router.get("/yetenek-matrisi")
async def jt_yetenek_matrisi(tenant: str = None, bolge: str = None, sube: str = None, rol: str = None):
    """Screen 5: 9-Box Talent Matrix"""
    q = _build_filter(tenant, bolge, sube, rol)
    karne = await db.jt_karne.find(q, {"_id": 0}).to_list(500)
    if not karne:
        return {"matrix": [], "high_potential": [], "risk_list": []}

    hedef_scores = sorted([k["hedef_skoru"] for k in karne])
    yetkinlik_scores = sorted([k["yetkinlik_skoru"] for k in karne])
    n = len(karne)
    h_low = hedef_scores[n // 3] if n > 2 else 70
    h_high = hedef_scores[2 * n // 3] if n > 2 else 90
    y_low = yetkinlik_scores[n // 3] if n > 2 else 70
    y_high = yetkinlik_scores[2 * n // 3] if n > 2 else 90

    def _box(h, y):
        hx = 0 if h < h_low else (1 if h < h_high else 2)
        yx = 0 if y < y_low else (1 if y < y_high else 2)
        return (hx, yx)

    LABELS = {
        (0,2): "Potansiyel Yetkinlik", (1,2): "Yükselen Yıldız", (2,2): "Yıldız",
        (0,1): "Tutarsız (Yetkinlik+)", (1,1): "Temel Katkı", (2,1): "Güçlü Performans",
        (0,0): "Gelişim Gerekli", (1,0): "Tutarsız (Hedef+)", (2,0): "Hedef Odaklı",
    }
    COLORS = {
        (0,2): "#3B82F6", (1,2): "#8B5CF6", (2,2): "#F59E0B",
        (0,1): "#6EE7B7", (1,1): "#94A3B8", (2,1): "#14B8A6",
        (0,0): "#EF4444", (1,0): "#FB923C", (2,0): "#60A5FA",
    }

    boxes = defaultdict(list)
    for k in karne:
        box = _box(k["hedef_skoru"], k["yetkinlik_skoru"])
        boxes[box].append({"sicil": k["sicil_no"], "ad": k["ad_soyad"], "sube": k["sube"],
                          "hedef": round(k["hedef_skoru"], 1), "yetkinlik": round(k["yetkinlik_skoru"], 1),
                          "toplam": round(k["toplam_skor"], 1), "segment": k["segment"]})

    matrix = []
    for (hx, yx), label in LABELS.items():
        people = boxes.get((hx, yx), [])
        matrix.append({"hedef_dilim": hx, "yetkinlik_dilim": yx, "label": label,
                       "count": len(people), "employees": people, "color": COLORS.get((hx, yx), "#94A3B8")})

    high_potential = boxes.get((2, 2), [])
    risk_list = boxes.get((0, 0), [])

    return {"matrix": matrix, "high_potential": high_potential, "risk_list": risk_list,
            "thresholds": {"hedef_low": round(h_low, 1), "hedef_high": round(h_high, 1),
                          "yetkinlik_low": round(y_low, 1), "yetkinlik_high": round(y_high, 1)}}


@router.get("/calisan-karnesi/{sicil_no}")
async def jt_calisan_karnesi(sicil_no: str, tenant: str = None):
    """Screen 6: Employee Report Card"""
    tq = {"tenant_id": tenant} if tenant else {}
    karne = await db.jt_karne.find_one({**tq, "sicil_no": sicil_no}, {"_id": 0})
    if not karne:
        raise HTTPException(404, "Çalışan karnesi bulunamadı")

    h_q = {**tq, "sicil_no": sicil_no}
    hedefler = await db.jt_hedefler.find(h_q, {"_id": 0}).to_list(20)
    yetkinlikler = await db.jt_yetkinlikler.find(h_q, {"_id": 0}).to_list(50)

    all_karne = await db.jt_karne.find(tq, {"_id": 0}).to_list(500)
    sube_peers = [k for k in all_karne if k["sube"] == karne["sube"] and k["sicil_no"] != sicil_no]
    bolge_peers = [k for k in all_karne if k["bolge"] == karne["bolge"] and k["sicil_no"] != sicil_no]
    sirket_avg = round(sum(k["toplam_skor"] for k in all_karne) / len(all_karne), 1) if all_karne else 0
    sube_avg = round(sum(k["toplam_skor"] for k in sube_peers) / len(sube_peers), 1) if sube_peers else 0
    bolge_avg = round(sum(k["toplam_skor"] for k in bolge_peers) / len(bolge_peers), 1) if bolge_peers else 0

    radar = [{"yetkinlik": y.get("yetkinlik", ""), "oz": y["oz_degerlendirme"], "yonetici": y["yonetici_puani"]} for y in yetkinlikler]

    return {
        "karne": karne,
        "hedefler": [{"hedef": h.get("hedef", ""), "agirlik": h["agirlik"], "birim": h.get("birim", ""),
                      "hedef_deger": h["hedef_deger"], "gerceklesen": h["gerceklesen"],
                      "gerceklesme": h["gerceklesme_pct"], "agirlikli_skor": h["agirlikli_skor"]} for h in hedefler],
        "yetkinlik_radar": radar,
        "karsilastirma": {"sirket": sirket_avg, "sube": sube_avg, "bolge": bolge_avg, "kisi": round(karne["toplam_skor"], 1)},
    }


@router.get("/filters")
async def jt_filters(tenant: str = None):
    """Return available filter values."""
    q = {"tenant_id": tenant} if tenant else {}
    karne = await db.jt_karne.find(q, {"_id": 0, "bolge": 1, "sube": 1, "rol": 1}).to_list(500)
    bolgeler = sorted(set(k.get("bolge", "") for k in karne if k.get("bolge")))
    subeler = sorted(set(k.get("sube", "") for k in karne if k.get("sube")))
    roller = sorted(set(k.get("rol", "") for k in karne if k.get("rol")))
    return {"bolgeler": bolgeler, "subeler": subeler, "roller": roller}


@router.get("/harita")
async def jt_harita(tenant: str = None):
    """Map data: region & branch level performance for Turkey map."""
    q = {"tenant_id": tenant} if tenant else {}
    karne = await db.jt_karne.find(q, {"_id": 0}).to_list(500)
    if not karne:
        return {"bolgeler": [], "subeler": []}

    # Branch → province mapping
    SUBE_IL = {
        "Kadıköy": "İstanbul", "Bakırköy": "İstanbul", "Nişantaşı": "İstanbul",
        "Genel Müdürlük": "İstanbul", "Çağrı Merkezi": "İstanbul",
        "Alsancak": "İzmir", "Kuşadası": "Aydın", "Bodrum": "Muğla",
        "Antalya Merkez": "Antalya", "Lara": "Antalya",
        "Ankara Kızılay": "Ankara", "Ankara Çankaya": "Ankara",
    }
    # Region → provinces mapping
    BOLGE_ILLER = {
        "Marmara": ["İstanbul", "Bursa", "Kocaeli", "Tekirdağ", "Balıkesir", "Çanakkale", "Edirne", "Kırklareli", "Sakarya", "Yalova", "Bilecik"],
        "Ege": ["İzmir", "Aydın", "Denizli", "Manisa", "Muğla", "Afyonkarahisar", "Kütahya", "Uşak"],
        "Akdeniz": ["Antalya", "Mersin", "Adana", "Hatay", "Burdur", "Isparta", "Kahramanmaraş", "Osmaniye"],
        "İç Anadolu": ["Ankara", "Konya", "Kayseri", "Eskişehir", "Sivas", "Yozgat", "Kırşehir", "Kırıkkale", "Aksaray", "Niğde", "Nevşehir", "Çankırı", "Karaman"],
        "Merkez": ["İstanbul"],
    }

    bolge_data = defaultdict(lambda: {"skorlar": [], "kisi": 0, "segments": Counter()})
    sube_data = defaultdict(lambda: {"skorlar": [], "kisi": 0, "segments": Counter()})
    for k in karne:
        b = k.get("bolge", "")
        s = k.get("sube", "")
        bolge_data[b]["skorlar"].append(k["toplam_skor"])
        bolge_data[b]["kisi"] += 1
        bolge_data[b]["segments"][k.get("segment", "")] += 1
        sube_data[s]["skorlar"].append(k["toplam_skor"])
        sube_data[s]["kisi"] += 1
        sube_data[s]["segments"][k.get("segment", "")] += 1

    bolge_result = []
    for b, d in bolge_data.items():
        avg = round(sum(d["skorlar"]) / len(d["skorlar"]), 1)
        bolge_result.append({
            "bolge": b, "iller": BOLGE_ILLER.get(b, []), "kisi": d["kisi"],
            "ort_skor": avg, "yildiz": d["segments"].get("Yıldız", 0),
            "segments": dict(d["segments"]),
        })

    sube_result = []
    for s, d in sube_data.items():
        avg = round(sum(d["skorlar"]) / len(d["skorlar"]), 1)
        sube_result.append({
            "sube": s, "il": SUBE_IL.get(s, ""), "kisi": d["kisi"],
            "ort_skor": avg, "yildiz": d["segments"].get("Yıldız", 0),
            "segments": dict(d["segments"]),
            "lat": _SUBE_COORDS.get(s, {}).get("lat"),
            "lng": _SUBE_COORDS.get(s, {}).get("lng"),
        })

    return {"bolgeler": bolge_result, "subeler": sube_result}


_SUBE_COORDS = {
    "Kadıköy": {"lat": 40.98, "lng": 29.03}, "Bakırköy": {"lat": 40.98, "lng": 28.87},
    "Nişantaşı": {"lat": 41.05, "lng": 28.99}, "Genel Müdürlük": {"lat": 41.01, "lng": 28.97},
    "Çağrı Merkezi": {"lat": 41.02, "lng": 29.01},
    "Alsancak": {"lat": 38.44, "lng": 27.14}, "Kuşadası": {"lat": 37.86, "lng": 27.26},
    "Bodrum": {"lat": 37.04, "lng": 27.43},
    "Antalya Merkez": {"lat": 36.89, "lng": 30.71}, "Lara": {"lat": 36.86, "lng": 30.76},
    "Ankara Kızılay": {"lat": 39.92, "lng": 32.85}, "Ankara Çankaya": {"lat": 39.90, "lng": 32.86},
}


@router.get("/calisan-listesi")
async def jt_calisan_listesi(tenant: str = None, bolge: str = None, sube: str = None, rol: str = None):
    """Return employee list for selection."""
    q = _build_filter(tenant, bolge, sube, rol)
    karne = await db.jt_karne.find(q, {"_id": 0}).sort("toplam_skor", -1).to_list(500)
    return {"calisanlar": [{"sicil": k["sicil_no"], "ad": k["ad_soyad"], "rol": k.get("rol", ""),
                           "sube": k.get("sube", ""), "toplam_skor": round(k["toplam_skor"], 1),
                           "segment": k.get("segment", "")} for k in karne]}
